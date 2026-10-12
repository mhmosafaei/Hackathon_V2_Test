"""Offline production-evidence regressions; original artifacts are read-only."""
from collections import Counter
from copy import deepcopy
from hashlib import sha256
import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from src.intelligence.contracts import CanonicalAnalysis
from src.intelligence.deterministic import build_projection, persist_projection
from src.permit.document_store import available_documents, load_document, save_decision
from src.permit.document_store import resolve_active_permit

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / 'cache/document_runs/94726c4cd85672f22deb19daededefa5fb72845b61bd18a60bbce4a1dbaba61e'


def saved():
    return json.loads((RUN / 'candidate_aggregate.json').read_text(encoding='utf-8'))


def test_kilkeel_projection_preserves_all_candidates_and_missing_conditions():
    original = saved()
    before = deepcopy(original)
    analysis, projection = build_projection(original, filename='permit.pdf', document_sha256='a'*64, analysis_id='local')
    assert original == before
    assert len(analysis.requirements) == len(projection['candidates']) == 37
    assert Counter(r.source_verification for r in analysis.requirements) == {'SOURCE_VERIFIED': 30, 'SOURCE_NOT_FOUND': 4, 'CONTEXT_UNRESOLVED': 3}
    groups = {g['condition_id']: g for g in projection['condition_groups']}
    assert set(groups) == {str(n) for n in range(1, 23)}
    assert groups['1']['status'] == groups['2']['status'] == 'NOT_CAPTURED'
    assert projection['coverage']['summary']['MISSING'] == 2
    assert analysis.coverage.completeness == 'PARTIAL'
    assert all(r.human_review == 'UNREVIEWED' for r in analysis.requirements)
    assert all(not r['executable'] for r in projection['evidence_mapping'])
    assert all(s.processor == 'DETERMINISTIC' for s in analysis.stage_runs[:3])
    assert analysis.audit.api_call_count == 0
    assert not analysis.constraint_proposals and not analysis.work_windows.evaluation_ready
    assert len(projection['duplicates']) == 3
    ais = [m for m in projection['evidence_mapping'] if m['condition_id'] in {'11', '12'}]
    assert len(ais) == 2 and all(m['evidence_source'] == 'AIS_GNSS' and m['automation_status'] == 'EXTERNAL_DATA_REQUIRED' for m in ais)


def test_page_19_correction_is_unique_source_backed_and_audited():
    import pymupdf
    with pymupdf.open(RUN / 'source.pdf') as pdf:
        page_texts = {i+1:p.get_text() for i,p in enumerate(pdf)}
    analysis, projection = build_projection(saved(), filename='permit.pdf', document_sha256='a'*64, analysis_id='local', page_texts=page_texts)
    row = next(r for r in analysis.requirements if r.parent_condition_id == '19')
    assert row.page_start == 3
    correction = next(c for c in projection['attribution_corrections'] if c['candidate_id'] == row.requirement_id)
    assert correction['original_page'] == 4 and correction['source_page'] == 3
    from src.documents.quote_verifier import normalise
    with pymupdf.open(RUN / 'source.pdf') as pdf:
        assert normalise(row.exact_quote) in normalise(pdf[2].get_text())
    assert next(c for c in saved()['candidates'] if c['candidate_id'] == row.requirement_id)['page_number'] == 4


def test_local_persistence_exports_registration_and_review_gate(tmp_path):
    folder = tmp_path / 'cache/document_runs/offline-projection'
    folder.mkdir(parents=True)
    original = (RUN / 'candidate_aggregate.json').read_bytes()
    (folder / 'candidate_aggregate.json').write_bytes(original)
    (folder / 'manifest.json').write_text(json.dumps({'document_state':'REVIEW_REQUIRED', 'analysis_run_id':'offline-projection'}))
    path = persist_projection(folder, filename='permit.pdf', document_sha256='a'*64)
    analysis = CanonicalAnalysis.model_validate_json(path.read_text())
    assert len(analysis.requirements) == 37
    assert (folder / 'candidate_aggregate.json').read_bytes() == original
    spec, rows, coverage = load_document(tmp_path, folder.name)
    assert len(rows) == 37 and spec['aggregate'].parent == path.parent
    assert len(coverage['condition_groups']) == 22
    book = load_workbook(io.BytesIO((path.parent / 'intelligence.xlsx').read_bytes()))
    assert book['Requirements'].max_row == 38
    assert book['Sources'].max_row == 42
    assert book['Condition groups'].max_row == 23
    assert book['Duplicates'].max_row == 4
    assert '\\u001c' in next(cell.value for row in book['Requirements'] for cell in row if isinstance(cell.value,str) and 'Dredged material' in cell.value)
    import csv
    assert len(list(csv.DictReader(io.StringIO((path.parent / 'requirements.csv').read_text())))) == 37
    with pytest.raises(ValueError, match='unverified'):
        save_decision(tmp_path, folder.name, next(c for c in rows if c['state']=='SOURCE_NOT_FOUND'), 'CONFIRMED', 'test')
    assert not (tmp_path/'cache/document_reviews').exists()
    first = path.read_bytes()
    assert persist_projection(folder, filename='permit.pdf', document_sha256='a'*64).read_bytes() == first


def test_ambiguous_source_page_is_not_silently_corrected():
    aggregate = saved()
    source = next(s for s in aggregate['source_inventory'] if s['section_reference'] == '19')
    duplicate = deepcopy(source)
    duplicate.update(source_record_id='ambiguous-source', pages=[6])
    aggregate['source_inventory'].append(duplicate)
    _, projected = build_projection(aggregate, filename='p.pdf', document_sha256='a'*64, analysis_id='test')
    assert not any(c['source_page'] == 6 for c in projected['attribution_corrections'])


def test_active_run_survives_restart_and_does_not_override_explicit_selection():
    documents = {'belfast': {}, 'fresh': {'analysis_run_id': 'run-fresh'}}
    assert resolve_active_permit(documents, None, 'run-fresh') == 'fresh'
    assert resolve_active_permit(documents, 'fresh', None) == 'fresh'
    assert resolve_active_permit(documents, 'belfast', 'run-fresh') == 'belfast'


def test_real_streamlit_cached_kilkeel_dashboard_renders_without_network(monkeypatch):
    from streamlit.testing.v1 import AppTest
    def no_network(*args, **kwargs):
        pytest.fail('Cached dashboard must not dispatch API requests')
    monkeypatch.setattr('src.documents.extractor.openai_candidate_request', no_network)
    monkeypatch.setattr('src.intelligence.service.openai_stage_transport', no_network)
    app = AppTest.from_file(str(ROOT/'app.py'), default_timeout=30)
    app.query_params['analysis_run_id'] = 'acceptance-b13ab39680004131b1a013dfd25505e1'
    app.run()
    assert not app.exception
    assert app.session_state['active_permit'] == RUN.name
    assert any('37 reviewable records' in c.value for c in app.caption)
    assert any('Source coverage gaps: 1 (NOT_CAPTURED), 2 (NOT_CAPTURED)' in w.value for w in app.warning)
