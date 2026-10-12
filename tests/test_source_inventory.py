from src.documents.source_inventory import build_source_inventory,link_candidates
from src.documents.pdf_parser import PdfPage
from src.documents.exports import excel_bytes
def test_inventory_links_clause_and_table_without_claiming_legal_meaning():
 pages=[PdfPage(1,'Conditions\n1.1.1 The operator shall report.\nTable S3.1 Limits')]
 inventory=build_source_inventory('d',pages)
 candidate={'page_number':1,'source_quotation':'1.1.1 The operator shall report.'}
 assert any(x.record_type=='CLAUSE' for x in inventory)
 assert link_candidates([candidate],inventory)[0]['source_record_ids']

def test_simple_numbered_conditions_are_inventoried_only_inside_conditions_section():
 pages=[PdfPage(1,'PART 1 - PARTICULARS\n1. Licensee Name'),
        PdfPage(2,'PART 2 - CONDITIONS\n1. The licensee shall notify the authority.\n2. Disposal must remain within Area A.'),
        PdfPage(3,'Appendix A Map of disposal area')]
 inventory=build_source_inventory('d',pages)
 clauses={record.section_reference for record in inventory if record.record_type=='CLAUSE'}
 assert clauses=={'1','2'}
 assert any(record.section_reference=='Part 1.1' for record in inventory)
 assert any('GEOGRAPHIC_EVIDENCE_ONLY' in record.ambiguity for record in inventory)
def test_excel_export_contains_real_workbook():
 assert excel_bytes({'document_id':'d','candidates':[],'source_inventory':[]})[:2]==b'PK'

def test_excel_export_has_reviewable_register_sheets_and_rows():
 from io import BytesIO
 from openpyxl import load_workbook
 payload=excel_bytes({'document_id':'d','run_id':'r','raw_candidate_count':1,
   'source_inventory':[{'record_id':'source-1','page_number':2,'source_text':'Condition text'}],
   'candidates':[{'candidate_id':'candidate-1','page_number':2,'state':'SOURCE_VERIFIED','human_review_required':True}],
   'coverage':{'processed_pages':[2]}})
 workbook=load_workbook(BytesIO(payload),read_only=True)
 assert workbook.sheetnames==['Document summary','Source inventory','Obligation register','Coverage','Review status']
 assert workbook['Source inventory'].max_row==2
 assert workbook['Obligation register'].max_row==2
 assert workbook['Review status'].max_row==2
 assert workbook['Coverage'].max_row==2

def test_inventory_links_explicit_cross_page_list_continuation():
 pages=[PdfPage(1,'4.1.1 Records shall:\n(a) be legible;\n(b) be retained'),PdfPage(2,'Permit number X\n2\n(c) be retrievable.\n4.1.2 Another requirement.')]
 inventory=build_source_inventory('d',pages)
 record=next(x for x in inventory if x.section_reference=='4.1.1')
 assert record.pages==[1,2]
 assert '(c) be retrievable.' in record.source_text

def test_table_row_keeps_only_explicitly_labelled_limit_context(monkeypatch):
 from src.documents import source_inventory
 monkeypatch.setattr(source_inventory,'reconstruct_tables',lambda _: [{'id':'t','page':1,'bbox':(0,0,1,1),'headers':['Parameter','Limit (including unit)'],'status':'VERIFIED_STRUCTURE','ambiguities':[],'rows':[[{'row':1,'column':0,'header':'Parameter','text':'TRO','bbox':None},{'row':1,'column':1,'header':'Limit (including unit)','text':'200 ug/l','bbox':None}]]}])
 row=next(x for x in build_source_inventory('d',[PdfPage(1,'Table')],'fake.pdf') if x.record_type=='TABLE_ROW')
 assert row.parameter=='TRO' and row.limit_text=='200 ug/l' and row.applicability is None
