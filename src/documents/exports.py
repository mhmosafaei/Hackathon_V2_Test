"""Review-safe exports sharing the exact persisted register records."""
from __future__ import annotations
import csv,io,json
import re
from pathlib import Path
def _excel_value(value):
 if isinstance(value,(list,dict)): return json.dumps(value)
 # XML cannot represent these source characters. Escape reversibly rather
 # than deleting evidence or changing persisted JSON/source quotations.
 if isinstance(value,str): return re.sub(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]',lambda m:f'\\u{ord(m.group()):04x}',value)
 return value
def csv_bytes(rows:list[dict])->bytes:
 fields=sorted({k for row in rows for k in row})
 out=io.StringIO(); writer=csv.DictWriter(out,fieldnames=fields,extrasaction='ignore'); writer.writeheader()
 for row in rows: writer.writerow({k:json.dumps(v) if isinstance(v,(list,dict)) else v for k,v in row.items()})
 return out.getvalue().encode()
def export_json(analysis:dict)->bytes: return json.dumps(analysis,indent=2).encode()
def excel_bytes(analysis:dict)->bytes:
 from openpyxl import Workbook
 from openpyxl.utils import get_column_letter
 wb=Workbook(); wb.remove(wb.active)
 coverage=analysis.get('coverage',{})
 coverage_rows=coverage.get('pages',[])
 if not coverage_rows and isinstance(coverage.get('processed_pages'),list):
  coverage_rows=[{'page':page,'status':'PROCESSED'} for page in coverage['processed_pages']]
 if analysis.get('schema_version') == 'permitpulse.intelligence.v1':
  sheets={
   'Document summary':[analysis.get('document',{})],
   'Sources':analysis.get('source_inventory',[]),
   'Requirements':analysis.get('requirements',[]),
   'Operational':analysis.get('operational_assessments',[]),
   'Constraints':analysis.get('constraint_proposals',[]),
   'Data Needs':analysis.get('data_needs',[]),
   'Data Feasibility':analysis.get('data_feasibility',[]),
   'Review Flags':analysis.get('review_flags',[]),
   'Condition groups':analysis.get('condition_groups',[]),
   'Evidence mapping':analysis.get('evidence_mapping',[]),
   'Attribution corrections':analysis.get('attribution_corrections',[]),
   'Duplicates':analysis.get('duplicates',[]),
   'Coverage':coverage.get('page_results',[]),
   'Work Windows':analysis.get('work_windows',{}).get('windows',[]),
   'Audit':analysis.get('stage_runs',[]) + [analysis.get('audit',{})],
  }
 else:
  sheets={'Document summary':[{'document_id':analysis.get('document_id'),'run_id':analysis.get('run_id'),'raw_candidates':analysis.get('raw_candidate_count')}], 'Source inventory':analysis.get('source_inventory',[]),'Obligation register':analysis.get('candidates',[]),'Coverage':coverage_rows,'Review status':[{'candidate_id':c['candidate_id'],'state':c['state'],'human_review_required':c['human_review_required']} for c in analysis.get('candidates',[])]}
 escaped_controls=False
 for name,rows in sheets.items():
  ws=wb.create_sheet(name); fields=sorted({k for row in rows if isinstance(row,dict) for k in row}) or ['value']; ws.append(fields)
  for row in rows:
   escaped_controls = escaped_controls or any(isinstance(row.get(k),str) and re.search(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]',row[k]) for k in fields)
   ws.append([_excel_value(row.get(k)) for k in fields])
   for cell in ws[ws.max_row]:
    if isinstance(cell.value,str): cell.data_type='s'
  ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
  for i in range(1,len(fields)+1): ws.column_dimensions[get_column_letter(i)].width=min(50,max(12,len(fields[i-1])+2))
 if escaped_controls:
  ws=wb.create_sheet('Export encoding'); ws.append(['Encoding note']); ws.append(['XML-illegal source control characters are displayed as literal \\uXXXX escapes. Canonical JSON preserves original characters. Text cells are not formulas.'])
 out=io.BytesIO(); wb.save(out); return out.getvalue()
